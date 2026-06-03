"""
BizDiag — Business Diagnostic Platform
FastAPI Backend
Supports: CSV, Excel, PDF, DOCX, TXT
Features: Full analysis engine, AI chat, export
"""

from fastapi import FastAPI, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from groq import Groq

import pandas as pd
import numpy as np
import io, os, re, traceback, json
from typing import Optional

# ── optional doc parsers (graceful fallback if missing) ──
try:
    import pdfplumber
    PDF_OK = True
except ImportError:
    PDF_OK = False

try:
    import docx as python_docx
    DOCX_OK = True
except ImportError:
    DOCX_OK = False

try:
    from openpyxl import load_workbook
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False

# ── export libs ──
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.units import cm
    REPORTLAB_OK = True
except ImportError:
    REPORTLAB_OK = False

load_dotenv()

GROQ_API_KEY = os.getenv("GROQ_API_KEY")
client = None
if GROQ_API_KEY:
    try:
        client = Groq(api_key=GROQ_API_KEY)
        print("✅ Groq Connected")
    except Exception as e:
        print("❌ Groq Error:", e)
else:
    print("⚠ No GROQ_API_KEY — AI disabled")

# =========================================================
# APP
# =========================================================

app = FastAPI(title="BizDiag API", version="9.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# =========================================================
# COLUMN ALIASES
# =========================================================

COLUMN_ALIASES = {
    "lead_source": ["lead_source","source","campaign","channel","traffic_source",
                    "منبع","کمپین","کانال","ورودی","سورس","منبع_لید"],
    "leads": ["leads","lead","lead_count","prospects","total_leads",
              "لید","سرنخ","تعداد_لید","تعداد لید"],
    "qualified_leads": ["qualified_leads","qualified","sql","mql",
                        "واجد_شرایط","لید_واجد"],
    "demos": ["demos","demo","calls","meetings","presentations",
              "جلسه","تماس","دمو","ارائه"],
    "proposals": ["proposals","proposal","quotes","offers",
                  "پیشنهاد","پروپوزال","پیشنهادیه"],
    "closed_deals": ["closed_deals","sales","orders","closed","deals","won",
                     "فروش","قرارداد","معامله","فروش_بسته"],
    "revenue": ["revenue","income","amount","total_revenue","sales_amount",
                "درآمد","مبلغ","سود","فروش_کل","کل_درآمد"],
    "marketing_cost": ["marketing_cost","cost","expense","budget","ad_cost",
                       "هزینه","بودجه","هزینه_تبلیغات","هزینه_بازاریابی"],
    "sales_rep": ["sales_rep","rep","agent","employee","salesperson",
                  "فروشنده","کارشناس","نماینده","کارمند"],
    "date": ["date","day","month","week","period",
             "تاریخ","روز","ماه","دوره"],
    "product": ["product","item","category","service","product_name",
                "محصول","کالا","دسته_بندی","خدمت"],
    "region": ["region","city","area","location","territory",
               "منطقه","شهر","ناحیه","استان"],
    "customer": ["customer","client","buyer","customer_name",
                 "مشتری","خریدار","کلاینت"],
    "deal_stage": ["deal_stage","stage","status","pipeline",
                   "مرحله","وضعیت","استیج"],
    "complaint": ["complaint","complaints","complaint_count","شکایت","تعداد_شکایت",
                  "issue","issues","problem","مشکل"],
    "satisfaction": ["satisfaction","nps","csat","rating","score","رضایت","امتیاز"],
    "churn": ["churn","churn_rate","attrition","ریزش","نرخ_ریزش"],
}

# =========================================================
# UTILS
# =========================================================

def normalize_text(text: str) -> str:
    if pd.isna(text): return ""
    text = str(text)
    text = text.replace("ي","ی").replace("ك","ک")
    for i, d in enumerate("۰۱۲۳۴۵۶۷۸۹"):
        text = text.replace(d, str(i))
    for ch in ["‌"," ","-","/","\\"]:
        text = text.replace(ch,"_")
    return text.strip().lower()

def smart_match_column(col_name: str) -> str | None:
    normalized = re.sub(r"[^a-z0-9آ-ی_]","", normalize_text(col_name))
    best_key, best_score = None, 0
    for key, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            alias_norm = re.sub(r"[^a-z0-9آ-ی_]","", normalize_text(alias))
            if normalized == alias_norm: return key
            if alias_norm in normalized or normalized in alias_norm:
                score = len(alias_norm)
                if score > best_score:
                    best_score, best_key = score, key
    return best_key if best_score >= 2 else None

def map_columns(df: pd.DataFrame) -> pd.DataFrame:
    rename = {}
    for col in df.columns:
        mapped = smart_match_column(col)
        if mapped and mapped not in rename.values():
            rename[col] = mapped
    return df.rename(columns=rename)

def safe_numeric(df: pd.DataFrame, cols: list) -> pd.DataFrame:
    for col in cols:
        if col not in df.columns:
            df[col] = 0
            continue
        df[col] = (
            df[col].astype(str)
            .str.replace(",","",regex=False).str.replace("٬","",regex=False)
            .str.replace("ریال","",regex=False).str.replace("تومان","",regex=False)
            .str.replace("$","",regex=False).str.replace("%","",regex=False)
            .str.replace(" ","",regex=False)
        )
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)
    return df

def safe_string_col(df: pd.DataFrame, col: str) -> pd.DataFrame:
    if col not in df.columns: df[col] = "Unknown"
    df[col] = df[col].fillna("Unknown").astype(str)
    return df

def safe_div(a, b, default=0):
    try: return a/b if b!=0 else default
    except: return default

def detect_encoding(raw: bytes) -> str:
    for enc in ["utf-8-sig","utf-8","cp1256","windows-1256","latin1"]:
        try:
            raw.decode(enc); return enc
        except: pass
    return "utf-8"

# =========================================================
# FILE READERS
# =========================================================

def read_csv_smart(raw: bytes) -> pd.DataFrame:
    if not raw: raise ValueError("Empty file")
    enc = detect_encoding(raw)
    best_df, best_cols = None, 0
    for e in [enc,"utf-8","utf-8-sig","cp1256","windows-1256","latin1"]:
        for sep in [",",";","\t","|"]:
            try:
                tmp = pd.read_csv(io.BytesIO(raw), encoding=e, sep=sep,
                                  engine="python", on_bad_lines="skip")
                if len(tmp.columns) > best_cols and len(tmp) > 0:
                    best_df, best_cols = tmp, len(tmp.columns)
            except: pass
    if best_df is None: raise ValueError("Could not parse CSV")
    return best_df

def read_excel_smart(raw: bytes) -> pd.DataFrame:
    xf = pd.ExcelFile(io.BytesIO(raw))
    best_df, best_rows = None, 0
    for sheet in xf.sheet_names:
        try:
            df = pd.read_excel(xf, sheet_name=sheet)
            if len(df) > best_rows:
                best_df, best_rows = df, len(df)
        except: pass
    if best_df is None: raise ValueError("No readable sheet")
    return best_df

def read_pdf_as_text(raw: bytes) -> str:
    """Extract text from PDF using pdfplumber."""
    if not PDF_OK:
        raise ValueError("pdfplumber not installed. Run: pip install pdfplumber")
    text_parts = []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t: text_parts.append(t)
    return "\n".join(text_parts)

def read_docx_as_text(raw: bytes) -> str:
    """Extract text from Word DOCX."""
    if not DOCX_OK:
        raise ValueError("python-docx not installed. Run: pip install python-docx")
    doc = python_docx.Document(io.BytesIO(raw))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def text_to_dataframe(text: str, filename: str) -> tuple[pd.DataFrame, str]:
    """Try to extract tabular data from raw text; return (df, raw_text)."""
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    # Try CSV-like parsing first
    for sep in [",","\t","|",";"]:
        try:
            df = pd.read_csv(io.StringIO(text), sep=sep, engine="python", on_bad_lines="skip")
            if len(df.columns) >= 3 and len(df) >= 2:
                return df, text
        except: pass
    # Return empty df with text for AI chat
    return pd.DataFrame(), text

# =========================================================
# ANALYSIS ENGINE
# =========================================================

def deep_analysis(df: pd.DataFrame) -> dict:
    result = {}
    num_cols = ["revenue","marketing_cost","leads","qualified_leads",
                "demos","proposals","closed_deals","complaint","satisfaction"]

    total_revenue   = float(df["revenue"].sum()) if "revenue" in df.columns else 0
    total_cost      = float(df["marketing_cost"].sum()) if "marketing_cost" in df.columns else 0
    total_leads     = int(df["leads"].sum()) if "leads" in df.columns else 0
    total_qualified = int(df["qualified_leads"].sum()) if "qualified_leads" in df.columns else 0
    total_demos     = int(df["demos"].sum()) if "demos" in df.columns else 0
    total_proposals = int(df["proposals"].sum()) if "proposals" in df.columns else 0
    total_closed    = int(df["closed_deals"].sum()) if "closed_deals" in df.columns else 0
    total_complaints= int(df["complaint"].sum()) if "complaint" in df.columns else 0

    conversion_rate   = safe_div(total_closed, total_leads) * 100
    win_rate          = safe_div(total_closed, total_proposals) * 100
    lead_quality_rate = safe_div(total_qualified, total_leads) * 100
    demo_rate         = safe_div(total_demos, total_qualified) * 100
    proposal_rate     = safe_div(total_proposals, total_demos) * 100
    roi               = safe_div(total_revenue - total_cost, total_cost) * 100
    cac               = safe_div(total_cost, total_closed)
    avg_deal          = safe_div(total_revenue, total_closed)
    profit_margin     = safe_div(total_revenue - total_cost, total_revenue) * 100
    roas              = safe_div(total_revenue, total_cost)
    complaint_rate    = safe_div(total_complaints, total_closed) * 100 if total_closed > 0 else 0

    result["kpis"] = {
        "total_revenue": round(total_revenue, 2),
        "marketing_cost": round(total_cost, 2),
        "total_leads": total_leads,
        "total_qualified_leads": total_qualified,
        "total_demos": total_demos,
        "total_proposals": total_proposals,
        "total_closed_deals": total_closed,
        "total_complaints": total_complaints,
        "conversion_rate": round(conversion_rate, 2),
        "win_rate": round(win_rate, 2),
        "lead_quality_rate": round(lead_quality_rate, 2),
        "demo_rate": round(demo_rate, 2),
        "proposal_rate": round(proposal_rate, 2),
        "roi": round(roi, 2),
        "cac": round(cac, 2),
        "average_deal_size": round(avg_deal, 2),
        "profit_margin": round(profit_margin, 2),
        "roas": round(roas, 2),
        "complaint_rate": round(complaint_rate, 2),
    }

    # ── funnel ──
    funnel = {
        "leads_to_qualified": round(safe_div(total_qualified, total_leads)*100, 1),
        "qualified_to_demo":  round(safe_div(total_demos, total_qualified)*100, 1),
        "demo_to_proposal":   round(safe_div(total_proposals, total_demos)*100, 1),
        "proposal_to_closed": round(safe_div(total_closed, total_proposals)*100, 1),
    }
    funnel_stages = [
        ("لیدها → واجد شرایط", funnel["leads_to_qualified"]),
        ("واجد شرایط → دمو", funnel["qualified_to_demo"]),
        ("دمو → پروپوزال", funnel["demo_to_proposal"]),
        ("پروپوزال → بسته‌شده", funnel["proposal_to_closed"]),
    ]
    weakest = min(funnel_stages, key=lambda x: x[1])
    result["funnel"] = funnel
    result["weakest_funnel_stage"] = weakest[0]
    result["weakest_funnel_pct"] = weakest[1]

    # ── monthly trend ──
    if "date" in df.columns:
        df["date"] = pd.to_datetime(df["date"], errors="coerce")
        df_d = df.dropna(subset=["date"])
        if len(df_d) > 0:
            monthly = (
                df_d.resample("ME", on="date")
                .agg(revenue=("revenue","sum"), leads=("leads","sum"),
                     closed=("closed_deals","sum"))
                .reset_index()
            )
            monthly["date"] = monthly["date"].dt.strftime("%Y-%m")
            result["monthly_trend"] = monthly.to_dict(orient="records")
            ts = (
                df_d.resample("W-SUN", on="date")
                .agg(revenue=("revenue","sum"), leads=("leads","sum"),
                     closed=("closed_deals","sum"))
                .reset_index()
            )
            ts["date"] = ts["date"].dt.strftime("%Y-%m-%d")
            result["time_series"] = ts.to_dict(orient="records")
            if len(monthly) >= 2:
                last  = float(monthly.iloc[-1]["revenue"])
                prev  = float(monthly.iloc[-2]["revenue"])
                result["mom_growth"] = round(safe_div(last - prev, prev)*100, 2)
            else:
                result["mom_growth"] = 0
        else:
            result["monthly_trend"] = []
            result["time_series"] = []
            result["mom_growth"] = 0
    else:
        result["monthly_trend"] = []
        result["time_series"] = []
        result["mom_growth"] = 0

    # ── source analysis ──
    if "lead_source" in df.columns:
        agg_dict = dict(
            revenue=("revenue", "sum"),
            leads=("leads", "sum"),
            closed=("closed_deals", "sum"),
        )
        if "marketing_cost" in df.columns:
            agg_dict["cost"] = ("marketing_cost", "sum")
        src = df.groupby("lead_source").agg(**agg_dict).reset_index()
        if "cost" not in src.columns:
            src["cost"] = 0
        src["conversion"] = (src["closed"]/src["leads"].replace(0,np.nan)*100).fillna(0).round(2)
        src["roi"] = ((src["revenue"]-src["cost"])/src["cost"].replace(0,np.nan)*100).fillna(0).round(2)
        src["roas"] = (src["revenue"]/src["cost"].replace(0,np.nan)).fillna(0).round(2)
        src = src.sort_values("revenue", ascending=False)
        result["source_analysis"] = src.to_dict(orient="records")
        result["best_lead_source"] = src.iloc[0]["lead_source"]
        result["worst_lead_source"] = src.iloc[-1]["lead_source"]
    else:
        result["source_analysis"] = []
        result["best_lead_source"] = "N/A"
        result["worst_lead_source"] = "N/A"

    # ── rep analysis ──
    if "sales_rep" in df.columns:
        rep_agg = dict(
            revenue=("revenue", "sum"),
            closed=("closed_deals", "sum"),
            leads=("leads", "sum"),
        )
        if "marketing_cost" in df.columns:
            rep_agg["cost"] = ("marketing_cost", "sum")
        rep = df.groupby("sales_rep").agg(**rep_agg).reset_index()
        if "cost" not in rep.columns:
            rep["cost"] = 0
        rep["avg_deal"] = (rep["revenue"]/rep["closed"].replace(0,np.nan)).fillna(0).round(2)
        rep["conversion"] = (rep["closed"]/rep["leads"].replace(0,np.nan)*100).fillna(0).round(2)
        rep = rep.sort_values("revenue", ascending=False)
        result["rep_analysis"] = rep.to_dict(orient="records")
        result["best_sales_rep"] = rep.iloc[0]["sales_rep"]
        result["worst_sales_rep"] = rep.iloc[-1]["sales_rep"]
    else:
        result["rep_analysis"] = []
        result["best_sales_rep"] = "N/A"
        result["worst_sales_rep"] = "N/A"

    # ── product analysis ──
    if "product" in df.columns:
        prod = (
            df.groupby("product")
            .agg(revenue=("revenue","sum"), closed=("closed_deals","sum"),
                 leads=("leads","sum"))
            .reset_index()
        )
        prod["avg_deal"] = (prod["revenue"]/prod["closed"].replace(0,np.nan)).fillna(0).round(2)
        prod = prod.sort_values("revenue", ascending=False)
        result["product_analysis"] = prod.to_dict(orient="records")
    else:
        result["product_analysis"] = []

    # ── region analysis ──
    if "region" in df.columns:
        reg = (
            df.groupby("region")
            .agg(revenue=("revenue","sum"), closed=("closed_deals","sum"))
            .reset_index()
            .sort_values("revenue", ascending=False)
        )
        result["region_analysis"] = reg.to_dict(orient="records")
    else:
        result["region_analysis"] = []

    # ── customer analysis ──
    customer_analysis = {}
    if "customer" in df.columns:
        cust = (
            df.groupby("customer")
            .agg(revenue=("revenue","sum"), closed=("closed_deals","sum"))
            .reset_index()
            .sort_values("revenue", ascending=False)
        )
        customer_analysis["top_customers"] = cust.head(10).to_dict(orient="records")
        customer_analysis["total_unique"] = int(df["customer"].nunique())
        customer_analysis["avg_revenue_per_customer"] = round(
            safe_div(total_revenue, df["customer"].nunique()), 2)
    result["customer_analysis"] = customer_analysis

    # ── complaint analysis ──
    complaint_analysis = {}
    if "complaint" in df.columns and df["complaint"].sum() > 0:
        complaint_analysis["total_complaints"] = total_complaints
        complaint_analysis["complaint_rate"] = round(complaint_rate, 2)
        complaint_analysis["complaint_severity"] = (
            "بحرانی 🔴" if complaint_rate > 20
            else "متوسط 🟡" if complaint_rate > 10
            else "پایین 🟢"
        )
        if "lead_source" in df.columns:
            comp_src = (
                df.groupby("lead_source")
                .agg(complaints=("complaint","sum"))
                .reset_index()
                .sort_values("complaints", ascending=False)
            )
            complaint_analysis["by_source"] = comp_src.to_dict(orient="records")
        if "product" in df.columns:
            comp_prod = (
                df.groupby("product")
                .agg(complaints=("complaint","sum"))
                .reset_index()
                .sort_values("complaints", ascending=False)
            )
            complaint_analysis["by_product"] = comp_prod.to_dict(orient="records")
    result["complaint_analysis"] = complaint_analysis

    # ── anomaly detection ──
    anomalies = []
    for col in ["revenue","leads","closed_deals"]:
        if col in df.columns:
            mean, std = df[col].mean(), df[col].std()
            if std > 0:
                high_out = df[df[col] > mean + 2*std]
                if len(high_out) > 0:
                    anomalies.append(f"ستون «{col}» دارای {len(high_out)} مقدار غیرعادی بالا است")
                low_out = df[df[col] < mean - 2*std]
                if len(low_out) > 0:
                    anomalies.append(f"ستون «{col}» دارای {len(low_out)} مقدار غیرعادی پایین است")
    result["anomalies"] = anomalies

    # ── business pathology (عارضه‌یابی) ──
    main_issues = []
    sub_issues  = []

    # Funnel issues
    if funnel["leads_to_qualified"] < 30:
        main_issues.append({
            "title": "کیفیت ناکافی لیدهای ورودی",
            "severity": "critical",
            "metric": f"{funnel['leads_to_qualified']:.1f}% واجد شرایط",
            "impact": "بیش از ۷۰٪ بودجه بازاریابی روی لیدهای بی‌کیفیت هدر می‌رود",
            "count": len(df)
        })
    if funnel["demo_to_proposal"] < 40:
        sub_issues.append({
            "title": "ضعف در تبدیل دمو به پروپوزال",
            "severity": "warning",
            "metric": f"{funnel['demo_to_proposal']:.1f}% تبدیل",
            "impact": "تیم فروش در متقاعد کردن مشتریان پس از دمو ضعیف عمل می‌کند"
        })
    if funnel["proposal_to_closed"] < 30:
        main_issues.append({
            "title": "نرخ پایین بستن قرارداد",
            "severity": "critical",
            "metric": f"{funnel['proposal_to_closed']:.1f}% بسته‌شده",
            "impact": "مشکل اساسی در قیمت‌گذاری، مذاکره یا ارائه ارزش به مشتری"
        })

    # ROI issues
    if roi < 0:
        main_issues.append({
            "title": "ROI منفی — زیان‌دهی بازاریابی",
            "severity": "critical",
            "metric": f"ROI: {roi:.1f}%",
            "impact": "هزینه‌های بازاریابی بیشتر از درآمد حاصل است — فوری اقدام کنید"
        })
    elif roi < 30:
        sub_issues.append({
            "title": "ROI پایین — بهره‌وری ناکافی",
            "severity": "warning",
            "metric": f"ROI: {roi:.1f}%",
            "impact": "بازگشت سرمایه کمتر از حد استاندارد صنعت"
        })

    # CAC issue
    if cac > avg_deal * 0.4 and avg_deal > 0:
        main_issues.append({
            "title": "هزینه جذب مشتری بیش از حد بالا",
            "severity": "critical",
            "metric": f"CAC: {cac:,.0f} vs میانگین معامله: {avg_deal:,.0f}",
            "impact": f"هزینه جذب هر مشتری {safe_div(cac,avg_deal)*100:.0f}٪ ارزش معامله را می‌بلعد"
        })

    # Growth issue
    if result["mom_growth"] < -15:
        main_issues.append({
            "title": "افت شدید درآمد ماه‌به‌ماه",
            "severity": "critical",
            "metric": f"رشد: {result['mom_growth']:+.1f}%",
            "impact": "روند نزولی درآمد — نیاز به بررسی فوری علل ریشه‌ای"
        })

    # Channel concentration risk
    if result["source_analysis"]:
        src_df = pd.DataFrame(result["source_analysis"])
        total_rev = src_df["revenue"].sum()
        top_share = safe_div(float(src_df.iloc[0]["revenue"]), float(total_rev)) * 100
        if top_share > 60:
            sub_issues.append({
                "title": "وابستگی بیش از حد به یک کانال",
                "severity": "warning",
                "metric": f"{src_df.iloc[0]['lead_source']}: {top_share:.0f}٪ درآمد",
                "impact": "ریسک تمرکز بالا — اگر این کانال ضعیف شود، کسب‌وکار در خطر است"
            })

    # Complaint issue
    if complaint_rate > 15:
        main_issues.append({
            "title": "نرخ شکایت بسیار بالا",
            "severity": "critical",
            "metric": f"{complaint_rate:.1f}٪ مشتریان شاکی",
            "impact": "خطر جدی برای اعتبار برند و نرخ حفظ مشتری"
        })

    # Win rate
    if win_rate < 20:
        main_issues.append({
            "title": "نرخ برد پایین در مذاکرات",
            "severity": "critical",
            "metric": f"Win Rate: {win_rate:.1f}%",
            "impact": "تیم فروش در بستن معاملات ضعیف عمل می‌کند — آموزش فوری لازم است"
        })

    # Rep performance gap
    if result["rep_analysis"] and len(result["rep_analysis"]) >= 2:
        reps_df = pd.DataFrame(result["rep_analysis"])
        best_rev = float(reps_df.iloc[0]["revenue"])
        worst_rev = float(reps_df.iloc[-1]["revenue"])
        gap = safe_div(best_rev - worst_rev, best_rev) * 100
        if gap > 50:
            sub_issues.append({
                "title": "شکاف عملکردی بزرگ بین فروشندگان",
                "severity": "warning",
                "metric": f"بهترین: {best_rev:,.0f} | ضعیف‌ترین: {worst_rev:,.0f}",
                "impact": f"شکاف {gap:.0f}٪ بین بهترین و ضعیف‌ترین فروشنده — استانداردسازی فرایند لازم است"
            })

    if not main_issues:
        main_issues.append({
            "title": "هیچ عارضه بحرانی شناسایی نشد",
            "severity": "healthy",
            "metric": "وضعیت سالم",
            "impact": "کسب‌وکار در وضعیت پایدار قرار دارد"
        })

    result["main_issues"] = main_issues
    result["sub_issues"] = sub_issues

    # Issue occurrence counts for chart
    issue_chart_data = []
    for issue in main_issues:
        issue_chart_data.append({
            "name": issue["title"][:30] + ("..." if len(issue["title"]) > 30 else ""),
            "count": issue.get("count", 1),
            "severity": issue["severity"]
        })
    for issue in sub_issues:
        issue_chart_data.append({
            "name": issue["title"][:30] + ("..." if len(issue["title"]) > 30 else ""),
            "count": issue.get("count", 1),
            "severity": issue["severity"]
        })
    result["issue_chart_data"] = issue_chart_data

    # ── key insights ──
    insights = []
    insights.append(f"📊 بهترین کانال جذب: {result['best_lead_source']} — بیشترین درآمد")
    insights.append(f"🏆 بهترین فروشنده: {result['best_sales_rep']}")
    insights.append(f"📉 ضعیف‌ترین مرحله قیف: {result['weakest_funnel_stage']} ({result['weakest_funnel_pct']:.1f}٪)")
    insights.append(f"💵 میانگین ارزش معامله: {avg_deal:,.0f}")
    insights.append(f"🔄 ROAS: {roas:.2f}x — به ازای هر ۱ واحد هزینه، {roas:.2f} واحد درآمد")
    insights.append(f"📈 رشد ماه‌به‌ماه: {result['mom_growth']:+.1f}٪")
    if total_complaints > 0:
        insights.append(f"⚠️ نرخ شکایت: {complaint_rate:.1f}٪ — {'نیاز به توجه' if complaint_rate > 10 else 'قابل قبول'}")
    if anomalies:
        insights.extend(anomalies[:2])
    result["insights"] = insights

    # ── problems ──
    problems = []
    if conversion_rate < 5: problems.append(f"🔴 نرخ تبدیل بسیار پایین ({conversion_rate:.1f}%) — هدف: ≥۱۰٪")
    elif conversion_rate < 10: problems.append(f"🟡 نرخ تبدیل پایین ({conversion_rate:.1f}%) — جای بهبود دارد")
    if roi < 0: problems.append(f"🔴 ROI منفی ({roi:.1f}%) — زیان‌دهی")
    elif roi < 20: problems.append(f"🟡 ROI پایین ({roi:.1f}%)")
    if win_rate < 20: problems.append(f"🔴 نرخ برد پایین ({win_rate:.1f}%)")
    if lead_quality_rate < 30: problems.append(f"🟡 کیفیت لید پایین ({lead_quality_rate:.1f}%)")
    if cac > avg_deal * 0.5 and avg_deal > 0: problems.append(f"🔴 CAC بیش از حد بالا ({cac:.0f} vs {avg_deal:.0f})")
    if result["mom_growth"] < -10: problems.append(f"🔴 افت درآمد: {result['mom_growth']:.1f}٪")
    result["problems"] = problems if problems else ["✅ هیچ مشکل بحرانی شناسایی نشد"]

    # ── recommendations ──
    short_term, mid_term, long_term = [], [], []
    if conversion_rate < 10:
        short_term.append("🎯 بازنگری فوری قیف فروش و شناسایی نقاط ریزش")
    if result["worst_lead_source"] != "N/A":
        short_term.append(f"📉 توقف یا کاهش بودجه کانال «{result['worst_lead_source']}»")
    if result["best_lead_source"] != "N/A":
        short_term.append(f"🚀 افزایش سرمایه‌گذاری در کانال «{result['best_lead_source']}»")
    if win_rate < 30:
        mid_term.append("🤝 برگزاری دوره آموزشی بستن معامله برای تیم فروش")
    if lead_quality_rate < 40:
        mid_term.append("🎯 طراحی کمپین‌های هدفمند برای جذب لیدهای با کیفیت‌تر")
    if cac > avg_deal * 0.3:
        mid_term.append("⚙️ بهینه‌سازی فرایند جذب مشتری برای کاهش CAC")
    if roi > 100:
        long_term.append(f"💰 ROI عالی ({roi:.0f}٪) — برنامه‌ریزی برای مقیاس‌پذیری")
    long_term.append("📊 پیاده‌سازی سیستم CRM برای پایش لحظه‌ای قیف فروش")
    long_term.append("🌱 ایجاد برنامه وفاداری مشتری برای کاهش نرخ ریزش")
    if not short_term: short_term = ["✅ استراتژی فعلی را ادامه دهید"]
    if not mid_term: mid_term = ["✅ بهینه‌سازی‌های تدریجی کافی است"]

    result["recommendations"] = {
        "short_term": short_term,
        "mid_term": mid_term,
        "long_term": long_term
    }

    # ── final verdict ──
    health = min(
        round(
            min(roi/150*35, 35) +
            min(conversion_rate/15*35, 35) +
            min(win_rate/50*30, 30),
            100
        ), 100
    )
    result["health_score"] = max(health, 0)
    result["final_verdict"] = (
        "🟢 کسب‌وکار در وضعیت عالی قرار دارد. ادامه استراتژی فعلی با تمرکز بر مقیاس‌پذیری."
        if health >= 70 else
        "🟡 کسب‌وکار در وضعیت متوسط. نیاز به بهینه‌سازی در چند حوزه کلیدی."
        if health >= 45 else
        "🔴 کسب‌وکار نیاز فوری به تحول دارد. عارضه‌های بحرانی شناسایی شده است."
    )

    return result

# =========================================================
# AI REPORT
# =========================================================

def generate_ai_report(analysis: dict, raw_text: str = "") -> str:
    if client is None:
        return "⚠ گزارش AI غیرفعال است. GROQ_API_KEY را در .env تنظیم کنید."

    kpis = analysis["kpis"]
    issues_text = "\n".join(
        f"  [{i['severity'].upper()}] {i['title']}: {i['metric']}"
        for i in analysis.get("main_issues", [])
    )
    sub_text = "\n".join(
        f"  {i['title']}: {i['metric']}"
        for i in analysis.get("sub_issues", [])
    )
    recs = analysis.get("recommendations", {})
    recs_text = (
        "کوتاه‌مدت: " + " | ".join(recs.get("short_term", [])) + "\n" +
        "میان‌مدت: " + " | ".join(recs.get("mid_term", [])) + "\n" +
        "بلندمدت: " + " | ".join(recs.get("long_term", []))
    )

    prompt = f"""
تو یک مشاور ارشد عارضه‌یابی کسب‌وکار با ۲۰ سال تجربه در بهینه‌سازی فروش و بازاریابی هستی.
این داده‌ها را تحلیل کن و گزارش اجرایی جامع بنویس.

═══ KPI های کلیدی ═══
💰 کل درآمد: {kpis['total_revenue']:,.0f}
💸 هزینه بازاریابی: {kpis['marketing_cost']:,.0f}
📈 ROI: {kpis['roi']:.1f}% | ROAS: {kpis['roas']:.2f}x
💵 میانگین معامله: {kpis['average_deal_size']:,.0f} | CAC: {kpis['cac']:,.0f}
🔄 حاشیه سود: {kpis['profit_margin']:.1f}%

═══ قیف فروش ═══
لیدها: {kpis['total_leads']:,} → واجد: {kpis['total_qualified_leads']:,} → دمو: {kpis['total_demos']:,} → پروپوزال: {kpis['total_proposals']:,} → بسته: {kpis['total_closed_deals']:,}
نرخ تبدیل: {kpis['conversion_rate']:.1f}% | نرخ برد: {kpis['win_rate']:.1f}%
ضعیف‌ترین مرحله: {analysis.get('weakest_funnel_stage','N/A')}

═══ عارضه‌های اصلی ═══
{issues_text}

═══ عارضه‌های فرعی ═══
{sub_text}

═══ پیشنهادها ═══
{recs_text}

═══ رشد ═══
رشد ماه‌به‌ماه: {analysis.get('mom_growth',0):+.1f}%
امتیاز سلامت کسب‌وکار: {analysis.get('health_score',0):.0f}/100

وظیفه: یک گزارش اجرایی ۶۰۰-۸۰۰ کلمه‌ای به فارسی بنویس شامل:

۱. **جمع‌بندی مدیریتی** (۱۰۰ کلمه) — وضعیت کلی با اعداد دقیق
۲. **تحلیل جامع عارضه‌های کسب‌وکار** — ریشه‌یابی هر مشکل و تاثیر آن
۳. **تحلیل قیف فروش** — نقاط دقیق ریزش مشتری
۴. **تحلیل عملکرد کانال‌ها** — کجا سرمایه هدر می‌رود
۵. **هشدارها و ریسک‌های بحرانی** — نیاز به اقدام فوری
۶. **برنامه عمل اولویت‌دار** — ۵ اقدام با بیشترین تاثیر

لحن: حرفه‌ای، صریح، مبتنی بر داده، با ذکر اعداد دقیق.
"""
    try:
        resp = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[{"role":"user","content":prompt}],
            temperature=0.2, max_tokens=1800,
        )
        return resp.choices[0].message.content
    except Exception as e:
        try:
            resp = client.chat.completions.create(
                model="llama-3.1-8b-instant",
                messages=[{"role":"user","content":prompt}],
                temperature=0.2, max_tokens=1400,
            )
            return resp.choices[0].message.content
        except Exception as e2:
            return f"⚠ خطا در AI: {str(e2)}"

# =========================================================
# AI CHAT
# =========================================================

def ai_chat_response(question: str, analysis_context: str, raw_text: str = "") -> str:
    if client is None:
        return "⚠ AI غیرفعال است. GROQ_API_KEY تنظیم نشده."

    context = analysis_context
    if raw_text and len(raw_text) < 8000:
        context += f"\n\n═══ محتوای فایل ═══\n{raw_text[:6000]}"

    prompt = f"""
تو یک مشاور هوشمند کسب‌وکار و عارضه‌یاب هستی.
داده‌های زیر از تحلیل فایل کاربر است:

{context}

سوال کاربر: {question}

پاسخ دقیق، کاربردی و مختصر (حداکثر ۳۰۰ کلمه) به فارسی بده.
اگر سوال درباره داده‌های فایل نیست، به بهترین شکل کمک کن.
"""
    try:
        resp = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[{"role":"user","content":prompt}],
            temperature=0.3, max_tokens=800,
        )
        return resp.choices[0].message.content
    except Exception as e:
        return f"⚠ خطا: {str(e)}"

# =========================================================
# ROUTES
# =========================================================

@app.get("/")
def home(): return {"status":"running","version":"9.0.0","name":"BizDiag"}

@app.get("/health")
def health(): return {"status":"ok","groq_enabled":client is not None,"version":"9.0.0"}

@app.post("/analyze")
async def analyze_file(file: UploadFile = File(...)):
    try:
        if not file.filename:
            return {"success":False,"error":"No file uploaded"}

        fname = file.filename.lower()
        raw = await file.read()
        raw_text = ""
        df = pd.DataFrame()

        # ── read file by type ──
        if fname.endswith(".csv"):
            df = read_csv_smart(raw)
        elif fname.endswith((".xlsx",".xls")):
            df = read_excel_smart(raw)
        elif fname.endswith(".pdf"):
            raw_text = read_pdf_as_text(raw)
            df, raw_text = text_to_dataframe(raw_text, fname)
        elif fname.endswith((".docx",".doc")):
            raw_text = read_docx_as_text(raw)
            df, raw_text = text_to_dataframe(raw_text, fname)
        elif fname.endswith(".txt"):
            raw_text = raw.decode(detect_encoding(raw))
            df, raw_text = text_to_dataframe(raw_text, fname)
        else:
            return {"success":False,"error":"فرمت پشتیبانی نمی‌شود. CSV, Excel, PDF, DOCX, TXT آپلود کنید."}

        # If no tabular data, return text-only result
        if df.empty or len(df) < 2:
            ai_summary = ""
            if client and raw_text:
                ai_summary = ai_chat_response(
                    "این فایل را خلاصه کن و نکات کلیدی کسب‌وکاری آن را بگو.",
                    "", raw_text
                )
            return {
                "success": True,
                "mode": "text_only",
                "raw_text": raw_text[:5000],
                "ai_summary": ai_summary,
                "dataset_shape": {"rows":0,"columns":0},
                "detected_columns": [],
                "sample_data": [],
                "kpis": {}, "funnel": {}, "source_analysis": [],
                "rep_analysis": [], "time_series": [], "monthly_trend": [],
                "product_analysis": [], "region_analysis": [],
                "customer_analysis": {}, "complaint_analysis": {},
                "advanced_analysis": {
                    "main_issues":[],"sub_issues":[],"issue_chart_data":[],
                    "insights":[],"recommendations":{"short_term":[],"mid_term":[],"long_term":[]},
                    "problems":[],"health_score":0,"final_verdict":"داده کافی برای تحلیل وجود ندارد",
                    "mom_growth":0,"weakest_funnel_stage":"N/A","anomalies":[]
                },
                "ai_report": ai_summary,
            }

        # ── clean df ──
        df.columns = [re.sub(r"[^a-zA-Z0-9آ-ی_]","_",normalize_text(c)) for c in df.columns]
        df = df.loc[:, ~df.columns.duplicated()]
        df = map_columns(df)
        df = df.loc[:, ~df.columns.duplicated()]

        for col in ["lead_source","sales_rep"]:
            if col not in df.columns: df[col] = "Unknown"
        df = safe_string_col(df,"lead_source")
        df = safe_string_col(df,"sales_rep")

        num_cols = ["revenue","marketing_cost","leads","qualified_leads",
                    "demos","proposals","closed_deals","complaint","satisfaction","churn"]
        df = safe_numeric(df, num_cols)

        # Auto-fill funnel
        if df["leads"].sum() <= 0: df["leads"] = df["leads"].where(df["leads"] > 0, 1)
        if df["closed_deals"].sum() <= 0: df["closed_deals"] = (df["revenue"] > 0).astype(int)
        if df["qualified_leads"].sum() <= 0: df["qualified_leads"] = (df["leads"] * 0.6).astype(int)
        if df["demos"].sum() <= 0: df["demos"] = (df["qualified_leads"] * 0.5).astype(int)
        if df["proposals"].sum() <= 0: df["proposals"] = (df["demos"] * 0.7).astype(int)

        analysis = deep_analysis(df)
        ai_report = generate_ai_report(analysis, raw_text)

        # Serialize
        def serialize(obj):
            if isinstance(obj, (np.integer,)): return int(obj)
            if isinstance(obj, (np.floating,)):
                # np.nan and np.inf are not valid JSON — convert to None
                f = float(obj)
                return None if (f != f or f == float('inf') or f == float('-inf')) else f
            if isinstance(obj, (np.bool_,)): return bool(obj)
            if isinstance(obj, float):
                return None if (obj != obj or obj == float('inf') or obj == float('-inf')) else obj
            if isinstance(obj, dict): return {k: serialize(v) for k, v in obj.items()}
            if isinstance(obj, list): return [serialize(i) for i in obj]
            return obj

        return serialize({
            "success": True,
            "mode": "tabular",
            "raw_text": raw_text[:3000] if raw_text else "",
            "dataset_shape": {"rows":int(df.shape[0]),"columns":int(df.shape[1])},
            "detected_columns": list(df.columns),
            "sample_data": df.head(10).replace({np.nan:None}).to_dict(orient="records"),
            "kpis": analysis["kpis"],
            "funnel": analysis["funnel"],
            "monthly_trend": analysis["monthly_trend"],
            "source_analysis": analysis["source_analysis"],
            "rep_analysis": analysis["rep_analysis"],
            "time_series": analysis["time_series"],
            "product_analysis": analysis.get("product_analysis",[]),
            "region_analysis": analysis.get("region_analysis",[]),
            "customer_analysis": analysis.get("customer_analysis",{}),
            "complaint_analysis": analysis.get("complaint_analysis",{}),
            "advanced_analysis": {
                "best_lead_source": analysis["best_lead_source"],
                "worst_lead_source": analysis["worst_lead_source"],
                "best_sales_rep": analysis["best_sales_rep"],
                "worst_sales_rep": analysis["worst_sales_rep"],
                "weakest_funnel_stage": analysis["weakest_funnel_stage"],
                "mom_growth": analysis["mom_growth"],
                "anomalies": analysis["anomalies"],
                "insights": analysis["insights"],
                "recommendations": analysis["recommendations"],
                "problems": analysis["problems"],
                "main_issues": analysis["main_issues"],
                "sub_issues": analysis["sub_issues"],
                "issue_chart_data": analysis["issue_chart_data"],
                "health_score": analysis["health_score"],
                "final_verdict": analysis["final_verdict"],
            },
            "ai_report": ai_report,
        })

    except Exception as e:
        print(traceback.format_exc())
        return {"success":False,"error":str(e)}


@app.post("/chat")
async def chat_endpoint(
    question: str = Form(...),
    analysis_json: str = Form(default="{}"),
    raw_text: str = Form(default="")
):
    try:
        analysis = json.loads(analysis_json) if analysis_json else {}
        context = json.dumps(analysis, ensure_ascii=False, indent=2)[:4000]
        answer = ai_chat_response(question, context, raw_text)
        return {"success":True,"answer":answer}
    except Exception as e:
        return {"success":False,"answer":f"خطا: {str(e)}"}


@app.post("/export/excel")
async def export_excel(analysis_json: str = Form(...)):
    """Export full analysis as Excel."""
    try:
        data = json.loads(analysis_json)
        output = io.BytesIO()
        sheets_written = 0
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            # KPIs — always written (even if empty, to guarantee at least one sheet)
            kpis = data.get("kpis", {})
            kpi_df = pd.DataFrame([
                {"شاخص": k, "مقدار": v} for k, v in kpis.items()
            ])
            kpi_df.to_excel(writer, sheet_name="KPI های کلیدی", index=False)
            sheets_written += 1

            # Monthly trend
            if data.get("monthly_trend"):
                pd.DataFrame(data["monthly_trend"]).to_excel(
                    writer, sheet_name="روند ماهانه", index=False)
                sheets_written += 1

            # Source analysis
            if data.get("source_analysis"):
                pd.DataFrame(data["source_analysis"]).to_excel(
                    writer, sheet_name="تحلیل کانال‌ها", index=False)
                sheets_written += 1

            # Rep analysis
            if data.get("rep_analysis"):
                pd.DataFrame(data["rep_analysis"]).to_excel(
                    writer, sheet_name="تحلیل فروشندگان", index=False)
                sheets_written += 1

            # Product analysis
            if data.get("product_analysis"):
                pd.DataFrame(data["product_analysis"]).to_excel(
                    writer, sheet_name="تحلیل محصولات", index=False)
                sheets_written += 1

            # Issues
            adv = data.get("advanced_analysis", {})
            issues = adv.get("main_issues", []) + adv.get("sub_issues", [])
            if issues:
                pd.DataFrame(issues).to_excel(
                    writer, sheet_name="عارضه‌های کسب‌وکار", index=False)
                sheets_written += 1

            # Recommendations
            recs = adv.get("recommendations", {})
            rec_rows = []
            for term, items in recs.items():
                for item in items:
                    rec_rows.append({"بازه": term, "پیشنهاد": item})
            if rec_rows:
                pd.DataFrame(rec_rows).to_excel(
                    writer, sheet_name="پیشنهادهای اجرایی", index=False)
                sheets_written += 1

        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition":"attachment; filename=bizdiag_report.xlsx"}
        )
    except Exception as e:
        return {"success":False,"error":str(e)}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)